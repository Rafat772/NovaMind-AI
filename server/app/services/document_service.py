import os
import uuid
import logging
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional, Tuple
from bson import ObjectId

import pypdf
import docx

from app.core.config import settings
from app.core.database import db
from app.services.embedding_service import (
    generate_single_embedding,
    generate_batch_embeddings,
    cosine_similarity,
    get_gemini_client,
)

logger = logging.getLogger("novamind.documents")


def ensure_upload_dir(user_id: str) -> str:
    """Ensures user-specific upload directory exists and returns its path."""
    user_dir = os.path.join(settings.UPLOAD_DIR, str(user_id))
    os.makedirs(user_dir, exist_ok=True)
    return user_dir


def sanitize_filename(filename: str) -> str:
    """Extracts a clean filename without directory traversal characters."""
    clean = os.path.basename(filename)
    # Remove any dangerous characters
    clean = "".join(c for c in clean if c.isalnum() or c in "._- ")
    return clean or "unnamed_document"


def extract_text_from_pdf(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Extracts text page-by-page from a PDF file using pypdf.
    Returns: (list of dicts with {'page': int, 'text': str}, total_pages)
    """
    pages_data = []
    reader = pypdf.PdfReader(file_path)
    total_pages = len(reader.pages)

    for idx, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
            pages_data.append({
                "page": idx + 1,
                "text": page_text.strip(),
            })
        except Exception as err:
            logger.warning(f"Error extracting page {idx + 1} from {file_path}: {err}")
            pages_data.append({"page": idx + 1, "text": ""})

    return pages_data, total_pages


def extract_text_from_docx(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Extracts text from a DOCX file using python-docx.
    Returns: (list of dicts with {'page': int, 'text': str}, estimated_pages)
    """
    doc = docx.Document(file_path)
    full_paragraphs = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            full_paragraphs.append(text)

    # Also extract any table content
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_paragraphs.append(" | ".join(row_text))

    full_text = "\n\n".join(full_paragraphs)
    # Estimate pages (~3000 chars per page)
    estimated_pages = max(1, len(full_text) // 3000 + (1 if len(full_text) % 3000 > 0 else 0))
    return [{"page": 1, "text": full_text}], estimated_pages


def extract_text_from_txt(file_path: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Extracts plain text from TXT or MD files.
    """
    content = ""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", errors="ignore") as f:
            content = f.read()

    estimated_pages = max(1, len(content) // 3000 + (1 if len(content) % 3000 > 0 else 0))
    return [{"page": 1, "text": content.strip()}], estimated_pages


def chunk_text(
    pages_data: List[Dict[str, Any]],
    chunk_size: int = 1000,
    chunk_overlap: int = 200,
) -> List[Dict[str, Any]]:
    """
    Splits document text into overlapping chunks while preserving page numbers.
    """
    chunks = []
    chunk_index = 0

    for item in pages_data:
        page_num = item["page"]
        text = item["text"]
        if not text:
            continue

        # Simple sliding window chunker
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk_content = text[start:end].strip()

            if chunk_content:
                chunks.append({
                    "chunk_index": chunk_index,
                    "page_number": page_num,
                    "content": chunk_content,
                    "char_count": len(chunk_content),
                })
                chunk_index += 1

            if end >= text_len:
                break
            start += chunk_size - chunk_overlap

    return chunks


def format_doc_response(doc: dict) -> dict:
    """Maps MongoDB document fields to the DocumentResponse schema."""
    return {
        "id": str(doc["_id"]),
        "user_id": str(doc.get("user_id")),
        "filename": doc.get("filename", "Untitled"),
        "stored_filename": doc.get("stored_filename", ""),
        "file_size": doc.get("file_size", 0),
        "file_type": doc.get("file_type", ""),
        "status": doc.get("status", "uploaded"),
        "error_message": doc.get("error_message"),
        "page_count": doc.get("page_count", 0),
        "chunk_count": doc.get("chunk_count", 0),
        "total_characters": doc.get("total_characters", 0),
        "created_at": doc.get("created_at") or datetime.now(timezone.utc),
        "updated_at": doc.get("updated_at") or datetime.now(timezone.utc),
    }


async def save_and_create_document(
    user_id: str,
    original_filename: str,
    content: bytes,
    file_type: str,
) -> dict:
    """
    Saves file bytes to disk and creates an initial document record in MongoDB.
    """
    clean_name = sanitize_filename(original_filename)
    ext = os.path.splitext(clean_name)[1].lower()
    stored_filename = f"{uuid.uuid4().hex}{ext}"

    user_dir = ensure_upload_dir(user_id)
    file_path = os.path.join(user_dir, stored_filename)

    with open(file_path, "wb") as f:
        f.write(content)

    file_size = len(content)
    now = datetime.now(timezone.utc)

    doc_data = {
        "user_id": str(user_id),
        "filename": clean_name,
        "stored_filename": stored_filename,
        "file_path": file_path,
        "file_size": file_size,
        "file_type": file_type.replace(".", "").lower(),
        "status": "processing",
        "error_message": None,
        "page_count": 0,
        "chunk_count": 0,
        "total_characters": 0,
        "chunks": [],
        "created_at": now,
        "updated_at": now,
    }

    result = await db.documents.insert_one(doc_data)
    doc_data["_id"] = result.inserted_id
    return doc_data


async def process_document_content(doc_id: str, file_path: str, file_type: str) -> dict:
    """
    Performs text extraction, chunking, and updates the document status to 'ready' or 'failed'.
    """
    now = datetime.now(timezone.utc)
    try:
        ext = file_type.lower()
        if ext == "pdf":
            pages_data, page_count = extract_text_from_pdf(file_path)
        elif ext in ["docx", "doc"]:
            pages_data, page_count = extract_text_from_docx(file_path)
        elif ext in ["txt", "md"]:
            pages_data, page_count = extract_text_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        chunks = chunk_text(pages_data)
        total_chars = sum(c["char_count"] for c in chunks)

        # Phase 3: Generate vector embeddings for all chunks
        chunk_texts = [c["content"] for c in chunks]
        embeddings = generate_batch_embeddings(chunk_texts)
        for idx, emb in enumerate(embeddings):
            chunks[idx]["embedding"] = emb

        update_fields = {
            "status": "ready",
            "page_count": page_count,
            "chunk_count": len(chunks),
            "total_characters": total_chars,
            "chunks": chunks,
            "error_message": None,
            "updated_at": now,
        }

        await db.documents.update_one(
            {"_id": ObjectId(doc_id)},
            {"$set": update_fields},
        )

        updated_doc = await db.documents.find_one({"_id": ObjectId(doc_id)})
        return updated_doc

    except Exception as exc:
        logger.error(f"Failed to process document {doc_id}: {exc}", exc_info=True)
        update_fields = {
            "status": "failed",
            "error_message": str(exc),
            "updated_at": now,
        }
        await db.documents.update_one(
            {"_id": ObjectId(doc_id)},
            {"$set": update_fields},
        )
        updated_doc = await db.documents.find_one({"_id": ObjectId(doc_id)})
        return updated_doc


async def get_user_documents(user_id: str) -> List[dict]:
    """Fetches all documents for a given user, newest first."""
    cursor = db.documents.find(
        {"user_id": str(user_id)},
        {"chunks": 0},  # Exclude chunks list for lightweight list response
    ).sort("created_at", -1)

    docs = await cursor.to_list(length=1000)
    return [format_doc_response(d) for d in docs]


async def get_document_by_id(doc_id: str, user_id: str) -> Optional[dict]:
    """Fetches a single document by ID belonging to user."""
    try:
        oid = ObjectId(doc_id)
    except Exception:
        return None

    doc = await db.documents.find_one({"_id": oid, "user_id": str(user_id)})
    if not doc:
        return None

    formatted = format_doc_response(doc)
    formatted["chunks"] = doc.get("chunks", [])
    return formatted


async def delete_document(doc_id: str, user_id: str) -> bool:
    """Deletes document record and removes file from disk."""
    try:
        oid = ObjectId(doc_id)
    except Exception:
        return False

    doc = await db.documents.find_one({"_id": oid, "user_id": str(user_id)})
    if not doc:
        return False

    # Remove file from disk
    file_path = doc.get("file_path")
    if file_path and os.path.exists(file_path):
        try:
            os.remove(file_path)
        except OSError as e:
            logger.warning(f"Could not remove physical file {file_path}: {e}")

    # Delete from MongoDB
    await db.documents.delete_one({"_id": oid})
    return True


async def get_document_stats(user_id: str) -> dict:
    """Calculates total docs, total chunks, total storage used."""
    pipeline = [
        {"$match": {"user_id": str(user_id)}},
        {
            "$group": {
                "_id": None,
                "total_documents": {"$sum": 1},
                "total_chunks": {"$sum": "$chunk_count"},
                "total_size_bytes": {"$sum": "$file_size"},
                "ready_count": {
                    "$sum": {"$cond": [{"$eq": ["$status", "ready"]}, 1, 0]}
                },
            }
        },
    ]

    cursor = db.documents.aggregate(pipeline)
    results = await cursor.to_list(length=1)

    if results:
        res = results[0]
        return {
            "total_documents": res.get("total_documents", 0),
            "total_chunks": res.get("total_chunks", 0),
            "total_size_bytes": res.get("total_size_bytes", 0),
            "ready_count": res.get("ready_count", 0),
        }

    return {
        "total_documents": 0,
        "total_chunks": 0,
        "total_size_bytes": 0,
        "ready_count": 0,
    }


async def semantic_search_chunks(
    user_id: str,
    query: str,
    document_id: Optional[str] = None,
    top_k: int = 5,
) -> List[dict]:
    """
    Performs cosine similarity vector search across chunks in user's documents.
    """
    query_vector = generate_single_embedding(query)
    query_filter: Dict[str, Any] = {"user_id": str(user_id), "status": "ready"}

    if document_id:
        try:
            query_filter["_id"] = ObjectId(document_id)
        except Exception:
            return []

    cursor = db.documents.find(
        query_filter,
        {"filename": 1, "chunks": 1},
    )
    docs = await cursor.to_list(length=100)

    scored_results = []
    for doc in docs:
        doc_id_str = str(doc["_id"])
        filename = doc.get("filename", "Untitled")
        chunks = doc.get("chunks", [])

        chunks_to_update = False
        for chunk in chunks:
            chunk_embedding = chunk.get("embedding")
            if not chunk_embedding or len(chunk_embedding) != len(query_vector):
                # If chunk was uploaded prior to embeddings or dimension mismatch, compute with current model
                chunk_embedding = generate_single_embedding(chunk["content"])
                chunk["embedding"] = chunk_embedding
                chunks_to_update = True

            score = cosine_similarity(query_vector, chunk_embedding)
            scored_results.append({
                "document_id": doc_id_str,
                "filename": filename,
                "chunk_index": chunk.get("chunk_index", 0),
                "page_number": chunk.get("page_number"),
                "content": chunk.get("content", ""),
                "score": round(score, 4),
            })

        # Save backfilled embeddings if needed
        if chunks_to_update:
            try:
                await db.documents.update_one(
                    {"_id": doc["_id"]},
                    {"$set": {"chunks": chunks}},
                )
            except Exception as e:
                logger.warning(f"Could not backfill embeddings for doc {doc_id_str}: {e}")

    # Rank by highest similarity score
    scored_results.sort(key=lambda x: x["score"], reverse=True)
    return scored_results[:top_k]


async def reindex_document(doc_id: str, user_id: str) -> Optional[dict]:
    """
    Regenerates embeddings for all chunks in a given document.
    """
    try:
        oid = ObjectId(doc_id)
    except Exception:
        return None

    doc = await db.documents.find_one({"_id": oid, "user_id": str(user_id)})
    if not doc:
        return None

    chunks = doc.get("chunks", [])
    if chunks:
        chunk_texts = [c["content"] for c in chunks]
        embeddings = generate_batch_embeddings(chunk_texts)
        for idx, emb in enumerate(embeddings):
            chunks[idx]["embedding"] = emb

        now = datetime.now(timezone.utc)
        await db.documents.update_one(
            {"_id": oid},
            {"$set": {"chunks": chunks, "updated_at": now}},
        )

    return format_doc_response(doc)


async def compare_documents(
    doc_id_a: str,
    doc_id_b: str,
    user_id: str,
    focus: str = "comprehensive",
) -> dict:
    """
    Performs comparative semantic synthesis between two indexed documents.
    """
    try:
        oid_a = ObjectId(doc_id_a)
        oid_b = ObjectId(doc_id_b)
    except Exception:
        raise ValueError("Invalid document ID format.")

    doc_a = await db.documents.find_one({"_id": oid_a, "user_id": str(user_id)})
    doc_b = await db.documents.find_one({"_id": oid_b, "user_id": str(user_id)})

    if not doc_a or not doc_b:
        raise ValueError("One or both documents could not be found.")

    chunks_a = doc_a.get("chunks", [])
    chunks_b = doc_b.get("chunks", [])

    text_a = "\n\n".join([c["content"] for c in chunks_a[:4]])
    text_b = "\n\n".join([c["content"] for c in chunks_b[:4]])

    name_a = doc_a.get("filename", "Document A")
    name_b = doc_b.get("filename", "Document B")

    client = get_gemini_client()
    markdown_result = ""

    if client:
        prompt = f"""You are NovaMind AI, an elite academic research assistant.
Compare and contrast the following two indexed research documents with high analytical rigor.

Document 1: {name_a}
Excerpts:
{text_a[:2500]}

---

Document 2: {name_b}
Excerpts:
{text_b[:2500]}

Focus Aspect: {focus}

Instructions:
1. Format your response cleanly in GitHub Markdown.
2. Structure your comparative analysis into:
   - ## 🎯 Executive Comparison Overview
   - ## 🔍 Core Distinct Themes & Methodological Contrasts
   - ## 🤝 Common Ground & Shared Patterns
   - ## 📊 Comparative Breakdown Table (Aspect | {name_a} | {name_b})
   - ## 💡 Synthesis & Key Recommendations
3. Keep the tone academic, crisp, and objective.
"""
        for model_name in ["gemini-flash-lite-latest", "gemini-3.1-flash-lite", "gemini-3.8-flash", "gemini-3.5-flash", "gemini-flash-latest"]:
            try:
                resp = client.models.generate_content(
                    model=model_name,
                    contents=prompt,
                )
                if resp and resp.text:
                    markdown_result = resp.text
                    break
            except Exception as e:
                logger.warning(f"Compare model {model_name} failed: {e}")
                continue

    if not markdown_result:
        markdown_result = f"""# Comparative Synthesis: {name_a} vs {name_b}

## 🎯 Executive Comparison Overview
* **Document A**: `{name_a}` ({doc_a.get('file_type', '').upper()} • {len(chunks_a)} chunks • {doc_a.get('page_count', 1)} pages)
* **Document B**: `{name_b}` ({doc_b.get('file_type', '').upper()} • {len(chunks_b)} chunks • {doc_b.get('page_count', 1)} pages)

## 🔍 Key Literature Focus
* **{name_a}**: Contains key research excerpts regarding:
> {chunks_a[0]['content'][:200] if chunks_a else "No excerpt available."}...

* **{name_b}**: Highlights foundational aspects regarding:
> {chunks_b[0]['content'][:200] if chunks_b else "No excerpt available."}...

## 📊 Quick Contrast Table
| Dimension | {name_a} | {name_b} |
| :--- | :--- | :--- |
| **Format** | {doc_a.get('file_type', '').upper()} | {doc_b.get('file_type', '').upper()} |
| **Length** | {doc_a.get('page_count', 1)} pages | {doc_b.get('page_count', 1)} pages |
| **Chunk Density** | {len(chunks_a)} passages | {len(chunks_b)} passages |

*Note: For deeper multi-paragraph LLM comparative breakdown, ensure your Gemini API key is active.*
"""

    return {
        "document_a": {
            "id": str(doc_a["_id"]),
            "filename": name_a,
            "file_type": doc_a.get("file_type", "").upper(),
            "page_count": doc_a.get("page_count", 1),
            "chunk_count": len(chunks_a),
        },
        "document_b": {
            "id": str(doc_b["_id"]),
            "filename": name_b,
            "file_type": doc_b.get("file_type", "").upper(),
            "page_count": doc_b.get("page_count", 1),
            "chunk_count": len(chunks_b),
        },
        "focus": focus,
        "comparison_markdown": markdown_result,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }

